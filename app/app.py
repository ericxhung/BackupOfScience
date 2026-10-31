from datetime import datetime
import io
import json
import os
import re
import string
from flask import Flask, make_response, render_template, request, send_file
import openpyxl
import pandas as pd

app = Flask(__name__)

# 自動計算上一層目錄作為根目錄
CURRENT_DIR = os.path.dirname(os.path.abspath(__file__))
PARENT_DIR = os.path.dirname(CURRENT_DIR)
FOLDER_PATH = os.path.join(PARENT_DIR, "電阻")
SETTING_FILE_PATH = os.path.join(PARENT_DIR, "setting.json")


def parse_date_from_string(text):
  match = re.search(r"(\d{4})_(\d{1,2})_(\d{1,2})", text)
  if match:
    y, m, d = match.groups()
    try:
      dt = datetime(int(y), int(m), int(d))
      return dt.strftime("%Y-%m-%d")
    except ValueError:
      pass
  return None


def load_item_mapping(all_sheets):
  # 載入 setting.json，若無則預設排序
  if os.path.exists(SETTING_FILE_PATH):
    try:
      with open(SETTING_FILE_PATH, "r", encoding="utf-8") as sf:
        mapping = json.load(sf)
        if isinstance(mapping, dict) and len(mapping) > 0:
          return mapping
    except Exception:
      pass
  return {sname: [sname] for sname in sorted(list(all_sheets))}


def process_experiment_data(mode="RESISTOR"):
  file_list = []
  if os.path.exists(FOLDER_PATH):
    for root, dirs, files in os.walk(FOLDER_PATH):
      for file in files:
        if file.lower().endswith((".xlsx", ".xls", ".xlsm")):
          f_path = os.path.join(root, file)
          folder_name = os.path.basename(root)

          date_str = parse_date_from_string(file)
          if not date_str:
            date_str = parse_date_from_string(folder_name)
          if not date_str:
            date_str = "2026-01-01"

          file_list.append({"path": f_path, "filename": file, "date": date_str})

  if not file_list:
    return pd.DataFrame(), {}

  all_sheets = set()
  for f in file_list:
    try:
      wb = openpyxl.load_workbook(f["path"], read_only=True, data_only=True)
      for sname in wb.sheetnames:
        if (
            sname not in ["氣溫", "溫度"]
            and "plxdaq" not in sname.lower()
        ):
          all_sheets.add(sname)
      wb.close()
    except Exception:
      pass

  item_mapping = load_item_mapping(all_sheets)
  matrix_data = {}
  cols = list(string.ascii_uppercase[:16])
  rows = range(1, 31)

  if mode == "RESISTOR":
    pattern = re.compile(r"^=\d+\*")
  else:
    pattern = re.compile(r"^=AVERAGE\(", re.IGNORECASE)

  for f in file_list:
    f_path = f["path"]
    dt_obj = datetime.strptime(f["date"], "%Y-%m-%d")
    date_col_name = dt_obj.strftime("%Y/%m/%d")

    if date_col_name not in matrix_data:
      matrix_data[date_col_name] = {}

    try:
      wb_formula = openpyxl.load_workbook(f_path, data_only=False)
      wb_data = openpyxl.load_workbook(f_path, data_only=True)

      for final_item_name, raw_sheets in item_mapping.items():
        found_val = ""
        for sname in raw_sheets:
          if sname in wb_formula.sheetnames and sname in wb_data.sheetnames:
            sheet_f = wb_formula[sname]
            sheet_d = wb_data[sname]

            found_coord = None
            for col in cols:
              for row in rows:
                cell_ref = f"{col}{row}"
                val = sheet_f[cell_ref].value
                if val is not None and isinstance(val, str):
                  if pattern.match(val):
                    found_coord = cell_ref
                    break
              if found_coord:
                break

            if found_coord:
              cell_val = sheet_d[found_coord].value
              if cell_val is not None:
                try:
                  found_val = float(cell_val)
                  break  # 找到第一筆就採用
                except ValueError:
                  pass

        matrix_data[date_col_name][final_item_name] = found_val

      wb_formula.close()
      wb_data.close()
    except Exception as e:
      print(f"處理檔案失敗 {f['filename']}: {e}")

  df = pd.DataFrame(matrix_data)
  final_index_order = list(item_mapping.keys())
  for sname in final_index_order:
    if sname not in df.index:
      df.loc[sname] = ""

  df = df.reindex(index=final_index_order)
  df = df.reindex(sorted(df.columns), axis=1)
  df.index.name = "項目"
  df = df.reset_index()

  return df, item_mapping


@app.route("/")
def index():
  mode = request.args.get("mode", "RESISTOR")
  df, _ = process_experiment_data(mode)

  if df.empty:
    return render_template(
        "index.html",
        tables="<p class='text-red-500'>找不到有效的 Excel 檔案或資料。</p>",
        dates=[],
        items=[],
        chart_json={},
        mode=mode,
    )

  dates = [col for col in df.columns if col != "項目"]
  items = df["項目"].tolist()

  # 準備前端 Chart.js 所需的資料結構
  chart_datasets = []
  colors = [
      "#3b82f6",
      "#ef4444",
      "#10b981",
      "#f59e0b",
      "#8b5cf6",
      "#ec4899",
      "#06b6d4",
      "#84cc16",
  ]

  for idx, item in enumerate(items):
    row_data = df[df["項目"] == item].iloc[0]
    values = [
        row_data[d] if pd.notna(row_data[d]) and row_data[d] != "" else None
        for d in dates
    ]
    chart_datasets.append({
        "label": item,
        "data": values,
        "borderColor": colors[idx % len(colors)],
        "backgroundColor": colors[idx % len(colors)],
        "fill": False,
        "tension": 0.1,
    })

  chart_json = {"labels": dates, "datasets": chart_datasets}

  # 產生 HTML 表格
  html_table = df.to_html(
      classes=(
          "table-auto w-full border-collapse border border-gray-200 text-center"
          " text-sm"
      ),
      index=False,
      na_rep="-",
  )

  return render_template(
      "index.html",
      tables=html_table,
      dates=dates,
      items=items,
      chart_json=chart_json,
      mode=mode,
  )


# 匯出路由 (支援 Excel, CSV, Markdown)
@app.route("/export/<file_type>")
def export_file(file_type):
  mode = request.args.get("mode", "RESISTOR")
  df, _ = process_experiment_data(mode)

  if df.empty:
    return "無數據可匯出", 400

  if file_type == "csv":
    output = io.BytesIO()
    output.write(df.to_csv(index=False, encoding="utf-8-sig").encode("utf-8-sig"))
    output.seek(0)
    return send_file(
        output,
        mimetype="text/csv",
        as_attachment=True,
        download_name="experiment_data.csv",
    )

  elif file_type == "md":
    headers = list(df.columns)
    header_str = " | ".join(str(h) for h in headers)
    sep_str = " | ".join("---" for _ in headers)
    md_lines = [f"| {header_str} |", f"| {sep_str} |"]
    for _, row in df.iterrows():
      row_str = " | ".join(
          str(val) if pd.notna(val) and val != "" else "" for val in row
      )
      md_lines.append(f"| {row_str} |")

    output = io.BytesIO()
    output.write("\n".join(md_lines).encode("utf-8"))
    output.seek(0)
    return send_file(
        output,
        mimetype="text/markdown",
        as_attachment=True,
        download_name="experiment_data.md",
    )

  else:  #預設 xlsx
    output = io.BytesIO()
    with pd.ExcelWriter(output, engine="openpyxl") as writer:
      df.to_excel(writer, index=False, sheet_name="數據整理")
    output.seek(0)
    return send_file(
        output,
        mimetype=(
            "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
        ),
        as_attachment=True,
        download_name="experiment_data.xlsx",
    )


if __name__ == "__main__":
  app.run(debug=True, port=5000)