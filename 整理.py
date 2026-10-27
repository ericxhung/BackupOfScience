import json
import os
import re
import string
from datetime import datetime
import tkinter as tk
from tkinter import filedialog, messagebox, ttk
import openpyxl
from openpyxl.styles import Font
import pandas as pd


class ExcelProcessorApp(tk.Tk):

  def __init__(self):
    super().__init__()
    self.title("Excel 資料處理與列印整合工具 (B欄固定寬度版)")
    self.geometry("900x700")

    # 數據整理全域變數
    self.file_list = []
    self.item_mapping = {}
    self.mode_var = tk.StringVar(value="RESISTOR")

    self.container = ttk.Frame(self)
    self.container.pack(fill="both", expand=True)

    # 進入點：顯示系統主頁
    self.show_main_home()

  def clear_container(self):
    for widget in self.container.winfo_children():
      widget.destroy()

  # ==================== 0. 系統主頁 (選單) ====================
  def show_main_home(self):
    self.clear_container()

    lbl_title = ttk.Label(
        self.container, text="Excel 資料處理系統", font=("Arial", 18, "bold")
    )
    lbl_title.pack(pady=40)

    btn_frame = ttk.Frame(self.container)
    btn_frame.pack(pady=20)

    btn_print = ttk.Button(
        btn_frame,
        text="🖨️ 數據列印",
        width=25,
        command=self.show_print_mode_select,
    )
    btn_print.pack(pady=15, ipady=10)

    btn_organize = ttk.Button(
        btn_frame,
        text="📊 數據整理",
        width=25,
        command=self.show_step1_home,
    )
    btn_organize.pack(pady=15, ipady=10)

  # ==================== 數據列印功能區 ====================
  def show_print_mode_select(self):
    file_path = filedialog.askopenfilename(
        title="選取數據列印用的 Excel 檔案",
        filetypes=[("Excel 檔案", "*.xlsx *.xls *.xlsm"), ("所有檔案", "*.*")],
    )
    if not file_path:
      return

    filename = os.path.basename(file_path)
    folder_name = os.path.basename(os.path.dirname(file_path))

    # 解析日期
    date_str = self.parse_date_from_string(filename)
    if not date_str:
      date_str = self.parse_date_from_string(folder_name)
    if not date_str:
      user_choice = self.request_manual_date(filename, folder_name)
      if user_choice == "SKIP" or user_choice is None:
        return
      date_str = user_choice

    self.show_print_layout_page(file_path, date_str)

  def show_print_layout_page(self, file_path, date_str):
    self.clear_container()

    lbl_title = ttk.Label(
        self.container, text="數據列印設定", font=("Arial", 14, "bold")
    )
    lbl_title.pack(pady=10)

    # 1. 上方設定區域 (日期與氣象參數)
    top_frame = ttk.LabelFrame(
        self.container, text="基本資訊設定", padding=10
    )
    top_frame.pack(fill="x", padx=20, pady=5)

    # 上方左邊：日期
    ttk.Label(top_frame, text="日期 (YYYY MM DD):").grid(
        row=0, column=0, sticky="w", padx=5, pady=5
    )
    self.entry_print_date = ttk.Entry(top_frame, width=15)
    self.entry_print_date.grid(row=0, column=1, sticky="w", padx=5, pady=5)
    formatted_date_input = date_str.replace("-", " ")
    self.entry_print_date.insert(0, formatted_date_input)

    # 上方右邊：溫度、濕度、天氣狀態
    ttk.Label(top_frame, text="溫度 (°C):").grid(
        row=0, column=2, sticky="w", padx=(20, 5), pady=5
    )
    self.entry_temp = ttk.Entry(top_frame, width=10)
    self.entry_temp.grid(row=0, column=3, sticky="w", padx=5, pady=5)

    ttk.Label(top_frame, text="濕度 (%):").grid(
        row=0, column=4, sticky="w", padx=(10, 5), pady=5
    )
    self.entry_humidity = ttk.Entry(top_frame, width=10)
    self.entry_humidity.grid(row=0, column=5, sticky="w", padx=5, pady=5)

    ttk.Label(top_frame, text="天氣狀態:").grid(
        row=0, column=6, sticky="w", padx=(10, 5), pady=5
    )
    self.entry_weather = ttk.Entry(top_frame, width=12)
    self.entry_weather.grid(row=0, column=7, sticky="w", padx=5, pady=5)
    self.entry_weather.insert(0, "晴")

    # 讀取 Excel 分析工作表與自動解析溫濕度
    self.print_file_path = file_path
    sheet_names = []
    auto_temp = ""
    auto_humidity = ""

    try:
      wb_data = openpyxl.load_workbook(file_path, data_only=True)
      sheet_names = wb_data.sheetnames

      # 尋找 氣溫 或 溫度 工作表
      target_temp_sheet = None
      for s in sheet_names:
        if s in ["氣溫", "溫度"]:
          target_temp_sheet = wb_data[s]
          break

      if target_temp_sheet:
        cols = list(string.ascii_uppercase[:16])
        rows = range(1, 31)

        temp_pattern = re.compile(r"(\d+(?:\.\d+)?)\s*(?:C|℃)", re.IGNORECASE)

        for col_idx, col in enumerate(cols):
          for row in rows:
            cell_ref = f"{col}{row}"
            val = target_temp_sheet[cell_ref].value
            if val is not None:
              val_str = str(val).strip()

              # 尋找溫度並取整數部分
              if not auto_temp:
                t_match = temp_pattern.search(val_str)
                if t_match:
                  try:
                    auto_temp = str(int(float(t_match.group(1))))
                  except ValueError:
                    auto_temp = t_match.group(1)

                  # 找到溫度後，直接抓取右邊那一格作為濕度
                  if col_idx + 1 < len(cols):
                    right_col = cols[col_idx + 1]
                    right_val = target_temp_sheet[
                        f"{right_col}{row}"
                    ].value
                    if right_val is not None:
                      if isinstance(right_val, (int, float)):
                        if 0 <= right_val <= 1:
                          hum_val = right_val * 100
                          auto_humidity = (
                              str(int(hum_val))
                              if hum_val.is_integer()
                              else str(round(hum_val, 1))
                          )
                        else:
                          auto_humidity = str(right_val)
                      else:
                        right_val_str = str(right_val).strip()
                        try:
                          f_val = float(right_val_str)
                          if 0 <= f_val <= 1:
                            hum_val = f_val * 100
                            auto_humidity = (
                                str(int(hum_val))
                                if hum_val.is_integer()
                                else str(round(hum_val, 1))
                            )
                          else:
                            auto_humidity = str(f_val)
                        except ValueError:
                          h_match = re.search(
                              r"(\d+(?:\.\d+)?)\s*[%％]?", right_val_str
                          )
                          if h_match:
                            f_val = float(h_match.group(1))
                            if 0 <= f_val <= 1:
                              hum_val = f_val * 100
                              auto_humidity = (
                                  str(int(hum_val))
                                  if hum_val.is_integer()
                                  else str(round(hum_val, 1))
                              )
                            else:
                              auto_humidity = str(f_val)

            if auto_temp and auto_humidity:
              break
          if auto_temp and auto_humidity:
            break

      wb_data.close()
    except Exception as e:
      messagebox.showerror("錯誤", f"讀取 Excel 檔案失敗：\n{e}")

    if auto_temp:
      self.entry_temp.insert(0, auto_temp)
    if auto_humidity:
      self.entry_humidity.insert(0, auto_humidity)

    # 2. 中央列表：項目勾選清單
    list_frame = ttk.LabelFrame(
        self.container, text="選取列印項目 (工作表)", padding=10
    )
    list_frame.pack(fill="both", expand=True, padx=20, pady=10)

    canvas = tk.Canvas(list_frame)
    scrollbar = ttk.Scrollbar(
        list_frame, orient="vertical", command=canvas.yview
    )
    scroll_frame = ttk.Frame(canvas)

    scroll_frame.bind(
        "<Configure>",
        lambda e: canvas.configure(scrollregion=canvas.bbox("all")),
    )
    canvas.create_window((0, 0), window=scroll_frame, anchor="nw")
    canvas.configure(yscrollcommand=scrollbar.set)

    canvas.pack(side="left", fill="both", expand=True)
    scrollbar.pack(side="right", fill="y")

    self.print_sheet_vars = {}
    for sname in sheet_names:
      is_check = (
          sname not in ["氣溫", "溫度"] and "plxdaq" not in sname.lower()
      )
      var = tk.BooleanVar(value=is_check)
      chk = ttk.Checkbutton(scroll_frame, text=sname, variable=var)
      chk.pack(anchor="w", pady=2)
      self.print_sheet_vars[sname] = var

    # 3. 下方控制按鈕
    bottom_frame = ttk.Frame(self.container)
    bottom_frame.pack(fill="x", padx=20, pady=10)

    btn_back = ttk.Button(
        bottom_frame, text="返回主選單", command=self.show_main_home
    )
    btn_back.pack(side="left")

    btn_confirm = ttk.Button(
        bottom_frame, text="確定 (生成列印用 Excel)", command=self.generate_print_excel
    )
    btn_confirm.pack(side="right")

  def generate_print_excel(self):
    date_val = self.entry_print_date.get().strip()
    match = re.match(r"^(\d{4})\s+(\d{1,2})\s+(\d{1,2})$", date_val)
    if not match:
      messagebox.showerror(
          "格式錯誤", "日期格式不正確，請輸入 YYYY MM DD（例如：2026 10 15）"
      )
      return

    year, month, day = match.groups()
    temp_val = self.entry_temp.get().strip()
    try:
      temp_val = str(int(float(temp_val)))
    except ValueError:
      pass

    humidity_val = self.entry_humidity.get().strip()
    weather_val = self.entry_weather.get().strip()

    selected_sheets = [
        s for s, var in self.print_sheet_vars.items() if var.get()
    ]
    if not selected_sheets:
      messagebox.showerror("錯誤", "請至少勾選一個項目！")
      return

    try:
      wb_formula = openpyxl.load_workbook(
          self.print_file_path, data_only=False
      )
      wb_data = openpyxl.load_workbook(self.print_file_path, data_only=True)

      cols = list(string.ascii_uppercase[:16])
      rows = range(1, 31)

      pattern_resistor = re.compile(r"^=\d+\*")
      pattern_average = re.compile(r"^=AVERAGE\(", re.IGNORECASE)

      export_rows = []

      for sname in selected_sheets:
        rx_val = ""
        a0_val = ""

        if sname in wb_formula.sheetnames and sname in wb_data.sheetnames:
          sheet_f = wb_formula[sname]
          sheet_d = wb_data[sname]

          # 尋找 Rx (電阻)
          found_rx_coord = None
          for col in cols:
            for row in rows:
              cell_ref = f"{col}{row}"
              v = sheet_f[cell_ref].value
              if v is not None and isinstance(v, str):
                if pattern_resistor.match(v):
                  found_rx_coord = cell_ref
                  break
            if found_rx_coord:
              break

          if found_rx_coord and sheet_d[found_rx_coord].value is not None:
            rx_val = sheet_d[found_rx_coord].value

          # 尋找 A0 (平均值)
          found_a0_coord = None
          for col in cols:
            for row in rows:
              cell_ref = f"{col}{row}"
              v = sheet_f[cell_ref].value
              if v is not None and isinstance(v, str):
                if pattern_average.match(v):
                  found_a0_coord = cell_ref
                  break
            if found_a0_coord:
              break

          if found_a0_coord and sheet_d[found_a0_coord].value is not None:
            a0_val = sheet_d[found_a0_coord].value

        export_rows.append((sname, rx_val, a0_val))

      wb_formula.close()
      wb_data.close()

      out_wb = openpyxl.Workbook()
      out_sheet = out_wb.active
      out_sheet.title = "數據列印"

      # 自動計算星期幾（只顯示數字，例如：一）
      dt_obj = datetime(int(year), int(month), int(day))
      weekdays_zh = ["一", "二", "三", "四", "五", "六", "日"]
      weekday_str = weekdays_zh[dt_obj.weekday()]

      # 填入 A1 整合資訊
      out_sheet["A1"] = (
          f"{year}年 {int(month)}月 {int(day)}日 ({weekday_str})"
          f" {temp_val}℃ {humidity_val}% {weather_val}"
      )

      # 填入表格標頭 (從 A3 開始)
      out_sheet["A3"] = "項目"
      out_sheet["B3"] = "Rx"
      out_sheet["C3"] = "A0"

      current_row = 4
      for name, rx, a0 in export_rows:
        out_sheet[f"A{current_row}"] = name
        out_sheet[f"B{current_row}"] = rx
        out_sheet[f"C{current_row}"] = a0
        current_row += 1

      # 字體大小改為 26
      for row in out_sheet.iter_rows():
        for cell in row:
          current_font = cell.font
          bold = current_font.bold if current_font else False
          italic = current_font.italic if current_font else False
          cell.font = Font(name="Arial", size=26, bold=bold, italic=italic)

      # 調整縱列寬度（A欄忽略A1，B欄固定為27.14）
      for col in out_sheet.columns:
        col_letter = openpyxl.utils.get_column_letter(col[0].column)

        if col_letter == "B":
          out_sheet.column_dimensions[col_letter].width = 27.14
          continue

        max_len = 0
        for cell in col:
          if cell.value is not None:
            if col_letter == "A" and cell.row == 1:
              continue

            val_str = str(cell.value)
            length = sum(2 if ord(c) > 127 else 1 for c in val_str)
            max_len = max(max_len, length)
        out_sheet.column_dimensions[col_letter].width = max(
            int(max_len * 2.5) + 6, 18
        )

      save_path = filedialog.asksaveasfilename(
          defaultextension=".xlsx",
          filetypes=[("Excel 檔案", "*.xlsx")],
          title="選擇儲存結果位置與檔名",
      )

      if not save_path:
        return

      out_wb.save(save_path)
      messagebox.showinfo("成功", f"列印用 Excel 檔案已成功生成至：\n{save_path}")

      # 自動用 Excel 開啟剛剛儲存的檔案 (Windows 專用)
      try:
        os.startfile(save_path)
      except Exception as e:
        print(f"自動開啟檔案失敗: {e}")

    except Exception as e:
      messagebox.showerror("重大錯誤", f"生成列印檔案時發生錯誤：\n{e}")

  # ==================== 1. 數據整理 - 步驟一：選擇檔案 (Home) ====================
  def show_step1_home(self):
    self.clear_container()

    lbl_title = ttk.Label(
        self.container,
        text="【數據整理】步驟一：選擇模式與管理 Excel 檔案",
        font=("Arial", 14, "bold"),
    )
    lbl_title.pack(pady=10)

    mode_frame = ttk.LabelFrame(
        self.container, text="搜尋模式設定", padding=10
    )
    mode_frame.pack(fill="x", padx=20, pady=5)

    radio_resistor = ttk.Radiobutton(
        mode_frame,
        text="1. 電阻模式 (對應 '=[數字]*[其他]' 格式)",
        value="RESISTOR",
        variable=self.mode_var,
    )
    radio_resistor.pack(anchor="w", pady=2)

    radio_average = ttk.Radiobutton(
        mode_frame,
        text="2. 平均值模式 (對應 '=AVERAGE([其他]' 格式)",
        value="AVERAGE",
        variable=self.mode_var,
    )
    radio_average.pack(anchor="w", pady=2)

    btn_frame = ttk.Frame(self.container)
    btn_frame.pack(fill="x", padx=20, pady=5)

    btn_add_files = ttk.Button(
        btn_frame, text="加入 Excel 檔案", command=self.add_excel_files
    )
    btn_add_files.pack(side="left", padx=5)

    btn_add_folder = ttk.Button(
        btn_frame, text="從資料夾遞迴加入", command=self.add_excel_folder
    )
    btn_add_folder.pack(side="left", padx=5)

    list_frame = ttk.Frame(self.container)
    list_frame.pack(fill="both", expand=True, padx=20, pady=10)

    self.tree_step1 = ttk.Treeview(
        list_frame, columns=("Filename", "Date", "Path"), show="headings"
    )
    self.tree_step1.heading("Filename", text="檔案名稱")
    self.tree_step1.heading("Date", text="解析日期 (YYYY-MM-DD)")
    self.tree_step1.heading("Path", text="完整路徑")
    self.tree_step1.column("Filename", width=220)
    self.tree_step1.column("Date", width=150)
    self.tree_step1.column("Path", width=380)
    self.tree_step1.pack(side="left", fill="both", expand=True)

    scrollbar = ttk.Scrollbar(
        list_frame, orient="vertical", command=self.tree_step1.yview
    )
    scrollbar.pack(side="right", fill="y")
    self.tree_step1.configure(yscrollcommand=scrollbar.set)

    bottom_frame = ttk.Frame(self.container)
    bottom_frame.pack(fill="x", padx=20, pady=10)

    btn_home = ttk.Button(
        bottom_frame, text="返回主選單", command=self.show_main_home
    )
    btn_home.pack(side="left", padx=5)

    btn_remove = ttk.Button(
        bottom_frame, text="移除選取檔案 (X)", command=self.remove_selected_file
    )
    btn_remove.pack(side="left", padx=5)

    btn_next = ttk.Button(
        bottom_frame, text="下一步 (管理與排序項目)", command=self.go_to_step2
    )
    btn_next.pack(side="right")

    self.refresh_step1_ui()

  def refresh_step1_ui(self):
    for item in self.tree_step1.get_children():
      self.tree_step1.delete(item)
    for idx, f in enumerate(self.file_list):
      self.tree_step1.insert(
          "", "end", iid=str(idx), values=(f["filename"], f["date"], f["path"])
      )

  def parse_date_from_string(self, text):
    match = re.search(r"(\d{4})_(\d{1,2})_(\d{1,2})", text)
    if match:
      y, m, d = match.groups()
      try:
        dt = datetime(int(y), int(m), int(d))
        return dt.strftime("%Y-%m-%d")
      except ValueError:
        pass
    return None

  def request_manual_date(self, filename, folder_name):
    top = tk.Toplevel(self)
    top.title(f"無法辨識日期: {filename}")
    top.geometry("420x200")
    top.grab_set()

    ttk.Label(
        top,
        text=(
            f"無法從檔名 '{filename}' 或資料夾 '{folder_name}' 辨識日期。\n"
            "請手動輸入正確日期（格式：YYYY MM DD），或點擊「不加入」略過："
        ),
        padding=10,
        wraplength=400,
    ).pack()

    entry_date = ttk.Entry(top, width=25)
    entry_date.pack(pady=5)
    entry_date.insert(0, "2026 01 01")

    result_date = [None]

    def on_submit():
      val = entry_date.get().strip()
      match = re.match(r"^(\d{4})\s+(\d{1,2})\s+(\d{1,2})$", val)
      if match:
        y, m, d = match.groups()
        try:
          dt = datetime(int(y), int(m), int(d))
          result_date[0] = dt.strftime("%Y-%m-%d")
          top.destroy()
          return
        except ValueError:
          pass
      messagebox.showerror(
          "格式錯誤",
          "日期格式不正確，請重新輸入（例如：2026 10 15）",
          parent=top,
      )

    def on_cancel():
      result_date[0] = "SKIP"
      top.destroy()

    btn_frame = ttk.Frame(top)
    btn_frame.pack(pady=10)

    ttk.Button(btn_frame, text="確定加入", command=on_submit).pack(
        side="left", padx=5
    )
    ttk.Button(btn_frame, text="不加入", command=on_cancel).pack(
        side="left", padx=5
    )

    self.wait_window(top)
    return result_date[0]

  def add_files_logic(self, file_paths):
    for path in file_paths:
      if not path.lower().endswith((".xlsx", ".xls", ".xlsm")):
        continue
      filename = os.path.basename(path)
      folder_name = os.path.basename(os.path.dirname(path))

      date_str = self.parse_date_from_string(filename)
      if not date_str:
        date_str = self.parse_date_from_string(folder_name)
      if not date_str:
        user_choice = self.request_manual_date(filename, folder_name)
        if user_choice == "SKIP" or user_choice is None:
          continue
        date_str = user_choice

      self.file_list.append(
          {"path": path, "filename": filename, "date": date_str}
      )
    self.refresh_step1_ui()

  def add_excel_files(self):
    files = filedialog.askopenfilenames(
        title="選擇 Excel 檔案",
        filetypes=[
            ("Excel 檔案", "*.xlsx *.xls *.xlsm"),
            ("所有檔案", "*.*"),
        ],
    )
    if files:
      self.add_files_logic(files)

  def add_excel_folder(self):
    folder = filedialog.askdirectory(title="選擇包含 Excel 的資料夾")
    if folder:
      matched_files = []
      try:
        for root, dirs, filenames in os.walk(folder):
          for file in filenames:
            if file.lower().endswith((".xlsx", ".xls", ".xlsm")):
              matched_files.append(os.path.join(root, file))
      except Exception as e:
        messagebox.showerror("錯誤", f"讀取資料夾時發生錯誤：\n{e}")
        return

      if matched_files:
        self.add_files_logic(matched_files)
        messagebox.showinfo(
            "加入完成", f"已從資料夾中尋找到 {len(matched_files)} 個 Excel 檔案！"
        )
      else:
        messagebox.showwarning(
            "提示", "該資料夾及其子資料夾中找不到任何支援的 Excel 檔案。"
        )

  def remove_selected_file(self):
    selected = self.tree_step1.selection()
    if not selected:
      messagebox.showwarning("提示", "請先選取要移除的檔案！")
      return
    for item in selected:
      idx = int(item)
      self.file_list[idx] = None
    self.file_list = [f for f in self.file_list if f is not None]
    self.refresh_step1_ui()

  def go_to_step2(self):
    if not self.file_list:
      messagebox.showerror("錯誤", "請至少加入一個 Excel 檔案才能進行下一步！")
      return
    self.show_step2_management()

  # ==================== 2. 數據整理 - 步驟二：項目管理 (Management) ====================
  def show_step2_management(self):
    self.clear_container()

    lbl_title = ttk.Label(
        self.container,
        text="【數據整理】步驟二：管理項目順序、合併與刪除工作表",
        font=("Arial", 14, "bold"),
    )
    lbl_title.pack(pady=10)

    all_sheets = set()
    for f in self.file_list:
      try:
        wb = openpyxl.load_workbook(f["path"], read_only=True, data_only=True)
        for sname in wb.sheetnames:
          all_sheets.add(sname)
        wb.close()
      except Exception as e:
        messagebox.showwarning(
            "讀取警告", f"無法讀取檔案 {f['filename']}：\n{e}"
        )

    setting_file_path = os.path.join(os.getcwd(), "setting.json")
    loaded_from_setting = False

    if os.path.exists(setting_file_path):
      try:
        with open(setting_file_path, "r", encoding="utf-8") as sf:
          saved_mapping = json.load(sf)
          if isinstance(saved_mapping, dict) and len(saved_mapping) > 0:
            self.item_mapping = saved_mapping
            loaded_from_setting = True
      except Exception:
        loaded_from_setting = False

    if not loaded_from_setting or not self.item_mapping:
      self.item_mapping = {sname: [sname] for sname in sorted(list(all_sheets))}

    lbl_info = ttk.Label(
        self.container,
        text=(
            "在此可調整項目上下順序、合併項目，或刪除不需要的項目。\n"
            "(提示：設定會自動儲存至 setting.json)"
        ),
    )
    lbl_info.pack(anchor="w", padx=20, pady=5)

    main_mgmt_frame = ttk.Frame(self.container)
    main_mgmt_frame.pack(fill="both", expand=True, padx=20, pady=5)

    self.listbox_items = tk.Listbox(
        main_mgmt_frame, selectmode=tk.EXTENDED, font=("Arial", 11)
    )
    self.listbox_items.pack(side="left", fill="both", expand=True)

    scrollbar = ttk.Scrollbar(
        main_mgmt_frame, orient="vertical", command=self.listbox_items.yview
    )
    scrollbar.pack(side="left", fill="y")
    self.listbox_items.configure(yscrollcommand=scrollbar.set)

    self.refresh_listbox()

    btn_ctrl_frame = ttk.Frame(main_mgmt_frame)
    btn_ctrl_frame.pack(side="right", fill="y", padx=10)

    btn_up = ttk.Button(
        btn_ctrl_frame, text="⬆ 往上移動", command=self.move_item_up
    )
    btn_up.pack(fill="x", pady=5)

    btn_down = ttk.Button(
        btn_ctrl_frame, text="⬇ 往下移動", command=self.move_item_down
    )
    btn_down.pack(fill="x", pady=5)

    sep1 = ttk.Separator(btn_ctrl_frame, orient="horizontal")
    sep1.pack(fill="x", pady=10)

    btn_merge = ttk.Button(
        btn_ctrl_frame, text="🔗 合併選取項目", command=self.merge_selected_items
    )
    btn_merge.pack(fill="x", pady=5)

    btn_delete = ttk.Button(
        btn_ctrl_frame, text="❌ 刪除選取項目", command=self.delete_selected_items
    )
    btn_delete.pack(fill="x", pady=5)

    bottom_frame = ttk.Frame(self.container)
    bottom_frame.pack(fill="x", padx=20, pady=10)

    btn_back = ttk.Button(
        bottom_frame, text="上一步", command=self.show_step1_home
    )
    btn_back.pack(side="left")

    btn_next = ttk.Button(
        bottom_frame, text="下一步 (開始計算並輸出)", command=self.go_to_step3
    )
    btn_next.pack(side="right")

  def refresh_listbox(self):
    self.listbox_items.delete(0, tk.END)
    for key in self.item_mapping.keys():
      self.listbox_items.insert(tk.END, key)

  def move_item_up(self):
    selected_indices = self.listbox_items.curselection()
    if not selected_indices:
      messagebox.showwarning("提示", "請先選取要移動的項目！")
      return
    items = list(self.item_mapping.items())
    for idx in selected_indices:
      if idx > 0:
        items[idx], items[idx - 1] = items[idx - 1], items[idx]
        self.item_mapping = dict(items)
        self.refresh_listbox()
        self.listbox_items.selection_set(idx - 1)
        self.listbox_items.selection_clear(idx)
        break

  def move_item_down(self):
    selected_indices = self.listbox_items.curselection()
    if not selected_indices:
      messagebox.showwarning("提示", "請先選取要移動的項目！")
      return
    items = list(self.item_mapping.items())
    for idx in reversed(selected_indices):
      if idx < len(items) - 1:
        items[idx], items[idx + 1] = items[idx + 1], items[idx]
        self.item_mapping = dict(items)
        self.refresh_listbox()
        self.listbox_items.selection_set(idx + 1)
        self.listbox_items.selection_clear(idx)
        break

  def delete_selected_items(self):
    selected_indices = self.listbox_items.curselection()
    if not selected_indices:
      messagebox.showwarning("提示", "請先選取要刪除的項目！")
      return

    items_keys = list(self.item_mapping.keys())
    selected_names = [items_keys[i] for i in selected_indices]

    if len(selected_names) == len(self.item_mapping):
      if not messagebox.askyesno(
          "警告",
          "你正準備刪除所有的項目！確定要清空嗎？",
          icon="warning",
      ):
        return

    for name in selected_names:
      if name in self.item_mapping:
        del self.item_mapping[name]

    self.refresh_listbox()
    messagebox.showinfo("成功", "已成功刪除選取的項目。")

  def merge_selected_items(self):
    selected_indices = self.listbox_items.curselection()
    if len(selected_indices) < 2:
      messagebox.showwarning("提示", "請至少選取兩個以上的項目進行合併！")
      return

    items_keys = list(self.item_mapping.keys())
    selected_names = [items_keys[i] for i in selected_indices]

    top = tk.Toplevel(self)
    top.title("合併項目與選擇代表名稱")
    top.geometry("420x220")
    top.grab_set()

    ttk.Label(
        top,
        text=(
            f"準備合併以下項目：\n{', '.join(selected_names)}\n\n"
            "請選擇其中一個作為最終的代表名稱："
        ),
        padding=10,
        wraplength=400,
    ).pack()

    combo_var = tk.StringVar(value=selected_names[0])
    combo_names = ttk.Combobox(
        top,
        textvariable=combo_var,
        values=selected_names,
        state="readonly",
        width=30,
    )
    combo_names.pack(pady=10)

    def on_confirm():
      chosen_name = combo_var.get()
      if not chosen_name:
        messagebox.showerror("錯誤", "請選擇一個代表名稱！", parent=top)
        return

      combined_sheets = []
      for name in selected_names:
        combined_sheets.extend(self.item_mapping[name])

      new_mapping = {}
      for k, v in self.item_mapping.items():
        if k not in selected_names:
          new_mapping[k] = v
        elif k == selected_names[0]:
          new_mapping[chosen_name] = combined_sheets

      self.item_mapping = new_mapping
      self.refresh_listbox()
      top.destroy()
      messagebox.showinfo("成功", f"已成功將項目合併，代表名稱為：{chosen_name}")

    ttk.Button(top, text="確認合併", command=on_confirm).pack(pady=10)

  def go_to_step3(self):
    if not self.item_mapping:
      messagebox.showerror("錯誤", "目前沒有任何有效的項目可供輸出！")
      return
    self.show_step3_output()

  # ==================== 3. 數據整理 - 步驟三：計算與矩陣輸出 (Output) ====================
  def show_step3_output(self):
    self.clear_container()

    lbl_title = ttk.Label(
        self.container,
        text="【數據整理】步驟三：自動搜尋儲存格、計算與矩陣匯出",
        font=("Arial", 14, "bold"),
    )
    lbl_title.pack(pady=10)

    mode_text = (
        "電阻模式 (=[數字]*...)"
        if self.mode_var.get() == "RESISTOR"
        else "平均值模式 (=AVERAGE(...)"
    )

    txt_info = ttk.Label(
        self.container,
        text=(
            f"目前選擇模式：【{mode_text}】\n"
            "已完成所有設定與項目管理。\n"
            "點擊下方按鈕將開始運算、儲存設定檔並匯出 Excel。"
        ),
    )
    txt_info.pack(padx=20, pady=20)

    btn_run = ttk.Button(
        self.container, text="開始計算並儲存 Excel", command=self.process_and_export
    )
    btn_run.pack(pady=10)

    btn_back_frame = ttk.Frame(self.container)
    btn_back = ttk.Button(
        btn_back_frame, text="上一步", command=self.show_step2_management
    )
    btn_back.pack()
    btn_back_frame.pack(pady=5)

  def process_and_export(self):
    try:
      setting_file_path = os.path.join(os.getcwd(), "setting.json")
      try:
        with open(setting_file_path, "w", encoding="utf-8") as sf:
          json.dump(self.item_mapping, sf, ensure_ascii=False, indent=4)
      except Exception as se:
        print(f"儲存 setting.json 失敗: {se}")

      matrix_data = {}
      cols = list(string.ascii_uppercase[:16])
      rows = range(1, 31)

      if self.mode_var.get() == "RESISTOR":
        pattern = re.compile(r"^=\d+\*")
      else:
        pattern = re.compile(r"^=AVERAGE\(", re.IGNORECASE)

      conflict_reports = []

      for f in self.file_list:
        f_path = f["path"]
        dt_obj = datetime.strptime(f["date"], "%Y-%m-%d")
        date_col_name = dt_obj.strftime("%Y/%m/%d")

        if date_col_name not in matrix_data:
          matrix_data[date_col_name] = {}

        try:
          wb_formula = openpyxl.load_workbook(f_path, data_only=False)
          wb_data = openpyxl.load_workbook(f_path, data_only=True)

          for final_item_name, raw_sheets in self.item_mapping.items():
            item_final_val = ""
            found_values_count = 0
            collected_values = []

            for sname in raw_sheets:
              if sname in wb_formula.sheetnames and sname in wb_data.sheetnames:
                sheet_f = wb_formula[sname]
                sheet_d = wb_data[sname]

                found_coord = None
                for col in cols:
                  for row in rows:
                    cell_ref = f"{col}{row}"
                    val = sheet_f[cell_ref].value
                    if val is not None and isinstance(val, str):
                      if pattern.match(val):
                        found_coord = cell_ref
                        break
                  if found_coord:
                    break

                if found_coord:
                  cell_val = sheet_d[found_coord].value
                  if cell_val is not None:
                    try:
                      val_num = float(cell_val)
                      collected_values.append((sname, val_num))
                      found_values_count += 1
                    except ValueError:
                      pass

            if found_values_count > 0:
              chosen_val = collected_values[0][1]

              if found_values_count > 1:
                sub_sheets_str = ", ".join(
                    [item[0] for item in collected_values]
                )
                conflict_reports.append(
                    f"日期: {date_col_name} | 檔案: {f['filename']} | 項目"
                    f" '{final_item_name}' 偵測到多筆重複數據"
                    f" ({sub_sheets_str})，已依規則採用代表項目數據。"
                )

              matrix_data[date_col_name][final_item_name] = chosen_val
            else:
              matrix_data[date_col_name][final_item_name] = ""

          wb_formula.close()
          wb_data.close()

        except Exception as e:
          messagebox.showerror(
              "錯誤", f"處理檔案 {f['filename']} 時發生錯誤：\n{e}"
          )
          return

      df = pd.DataFrame(matrix_data)

      final_index_order = list(self.item_mapping.keys())
      for sname in final_index_order:
        if sname not in df.index:
          df.loc[sname] = ""

      df = df.reindex(index=final_index_order)
      df = df.reindex(sorted(df.columns), axis=1)

      df.index.name = "項目"
      df = df.reset_index()

      save_path = filedialog.asksaveasfilename(
          defaultextension=".xlsx",
          filetypes=[("Excel 檔案", "*.xlsx")],
          title="選擇儲存結果位置與檔名",
      )

      if not save_path:
        return

      df.to_excel(save_path, index=False)

      if conflict_reports:
        warning_msg = (
            "資料已成功計算與匯出，設定檔已自動儲存。\n"
            "但發現以下項目在寫入時有重複數據碰撞：\n\n"
            + "\n".join(conflict_reports[:10])
        )
        if len(conflict_reports) > 10:
          warning_msg += (
              f"\n...以及其他 {len(conflict_reports) - 10} 筆重複記錄。"
          )
        messagebox.showwarning("資料重複警告提示", warning_msg)
      else:
        messagebox.showinfo(
            "成功",
            "資料已成功排序、合併、輸出 (日期格式 YYYY/MM/DD)，\n"
            f"設定檔已儲存至：{setting_file_path}\n"
            f"儲存路徑：\n{save_path}",
        )

    except Exception as e:
      messagebox.showerror("重大錯誤", f"輸出過程發生未預期的錯誤：\n{e}")


if __name__ == "__main__":
  app = ExcelProcessorApp()
  app.mainloop()